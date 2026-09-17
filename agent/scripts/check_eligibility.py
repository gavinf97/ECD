#!/usr/bin/env python3
"""
check_eligibility.py — Checklist Section 1 pre-checks + applicant-detail hints.

Matches the resource (name + homepage domain) against references/elixir_reference_lists.json:
  S1.2  listed as an ELIXIR Node service on the SDP (Node pages)       -> sdp_match (+ Node labels)
  S1.4  already a CDR / EDD / RIR / GCBR                               -> status_match
  S1.1  Node & Node institute                                          -> node_match (from SDP listing;
                                                                          affiliations/institution countries as hints)
  S1.3  resource type (deposition / knowledgebase / registry / aggregation) -> resource_type_signals
  S1.5  federated / consortium signals                                 -> federation_signals
  APP.community                                                        -> community_suggestions (keyword heuristic)

Optionally merges an enriched ELIXIR services spreadsheet (--services-xlsx; sheet "Services",
header row 4, columns incl. "Resource name", "URL", "ELIXIR Node", "ELIXIR badge") as a second SDP source.

Usage:
    python check_eligibility.py --identity WORKDIR/identity.json --evidence-dir WORKDIR/evidence
                                --out WORKDIR/evidence/eligibility.json [--services-xlsx PATH]
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import emit, load_json, load_reference_lists, log, norm_name, now_iso, same_site  # noqa: E402

COUNTRY_TO_NODE = {
    "belgium": "ELIXIR-Belgium", "czech": "ELIXIR-Czech Republic", "denmark": "ELIXIR-Denmark", "estonia": "ELIXIR-Estonia",
    "finland": "ELIXIR-Finland", "france": "ELIXIR-France", "germany": "ELIXIR-Germany", "greece": "ELIXIR-Greece",
    "hungary": "ELIXIR-Hungary", "ireland": "ELIXIR-Ireland", "israel": "ELIXIR-Israel", "italy": "ELIXIR-Italy",
    "luxembourg": "ELIXIR-Luxembourg", "netherlands": "ELIXIR-Netherlands", "norway": "ELIXIR-Norway",
    "portugal": "ELIXIR-Portugal", "slovenia": "ELIXIR-Slovenia", "spain": "ELIXIR-Spain", "sweden": "ELIXIR-Sweden",
    "switzerland": "ELIXIR-Switzerland", "united kingdom": "ELIXIR-UK", " uk": "ELIXIR-UK", "cyprus": "ELIXIR-Cyprus",
    "slovakia": "ELIXIR-Slovakia", "embl-ebi": "EMBL-EBI", "european bioinformatics institute": "EMBL-EBI",
}
ISO3_TO_NODE = {"BEL": "ELIXIR-Belgium", "CZE": "ELIXIR-Czech Republic", "DNK": "ELIXIR-Denmark", "EST": "ELIXIR-Estonia",
                "FIN": "ELIXIR-Finland", "FRA": "ELIXIR-France", "DEU": "ELIXIR-Germany", "GRC": "ELIXIR-Greece",
                "HUN": "ELIXIR-Hungary", "IRL": "ELIXIR-Ireland", "ISR": "ELIXIR-Israel", "ITA": "ELIXIR-Italy",
                "LUX": "ELIXIR-Luxembourg", "NLD": "ELIXIR-Netherlands", "NOR": "ELIXIR-Norway", "PRT": "ELIXIR-Portugal",
                "SVN": "ELIXIR-Slovenia", "ESP": "ELIXIR-Spain", "SWE": "ELIXIR-Sweden", "CHE": "ELIXIR-Switzerland",
                "GBR": "ELIXIR-UK", "CYP": "ELIXIR-Cyprus", "SVK": "ELIXIR-Slovakia"}
# Heuristic keyword hints only — Claude must judge relevance (S2.1d) from the Community's actual scope.
COMMUNITY_KEYWORDS = {
    "3D-BioInfo": r"structur|3d|fold|protein structure|pdb|cryo|macromolecul",
    "Biodiversity": r"biodiversit|species|taxonom|occurrence|specimen|ecolog|barcod",
    "Cancer Data": r"cancer|tumou?r|oncolog|somatic",
    "Domestic Animals Genome and Phenome": r"livestock|cattle|pig|sheep|goat|chicken|domestic animal|farm animal",
    "Federated Human Data": r"human genom|clinical|patient|sensitive|cohort|biobank|phenotyp",
    "Food and Nutrition": r"food|nutrition|diet|nutrient",
    "Galaxy": r"galaxy|workflow",
    "Human Copy Number Variation": r"copy number|cnv|structural variant",
    "Intrinsically Disordered Proteins": r"disorder|idp|idr|low complexity|phase separation|condensate",
    "Metabolomics": r"metabolom|metabolite|mass spectrometry|nmr",
    "Microbial Biotechnology": r"microb|strain|biotechnolog|bacteri|fung|enzyme",
    "Microbiome": r"microbiome|metagenom|16s|amplicon|microbial communit",
    "Plant Sciences": r"plant|crop|arabidopsis|agricultur|wheat|rice|maize|seed",
    "Proteomics": r"proteom|peptide|mass spectrometry|post-translational",
    "Rare Diseases": r"rare disease|orphan|mendelian",
    "Research Data Management": r"data management|fair data|metadata|data steward",
    "Single-Cell Omics": r"single[- ]cell|scrna|cell atlas",
    "Systems Biology": r"systems biology|pathway|kinetic|sbml|network model|mathematical model",
    "Toxicology": r"toxic|adverse outcome|chemical safety|hazard",
}


def matches(item: dict, name: str, url: str, aliases: set[str]) -> list[str]:
    basis = []
    if item.get("url") and same_site(item["url"], url):
        basis.append("homepage_domain")
    iname = norm_name(re.sub(r"\(.*?\)", "", item.get("name", "")))
    full = norm_name(item.get("name", ""))
    if iname in aliases or full in aliases:
        basis.append("name")
    elif any(a and len(a) > 3 and re.search(rf"\b{re.escape(a)}\b", re.sub(r"[^a-z0-9 ]", "", item.get("name", "").lower().replace("-", ""))) for a in aliases):
        basis.append("name_token")
    return basis


def load_xlsx_services(path: str) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Services"] if "Services" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows[:10]) if r and "Resource name" in [str(c).strip() if c else "" for c in r])
    header = [str(c).strip() if c else "" for c in rows[header_idx]]
    out = []
    for r in rows[header_idx + 1:]:
        rec = dict(zip(header, r))
        if not rec.get("Resource name"):
            continue
        out.append({"name": str(rec.get("Resource name")), "url": rec.get("url") or rec.get("URL"),
                    "node": rec.get("ELIXIR Node"), "badge": rec.get("ELIXIR badge"),
                    "functional_division": rec.get("Functional division"), "biotools": rec.get("bio.tools link"),
                    "fairsharing": rec.get("FAIRsharing link"), "last_checked": rec.get("Last checked")})
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--identity", required=True)
    ap.add_argument("--evidence-dir", required=True)
    ap.add_argument("--services-xlsx")
    ap.add_argument("--out")
    args = ap.parse_args()
    identity = load_json(args.identity)
    ev = Path(args.evidence_dir)
    refs = load_reference_lists()
    name = identity["input"]["name"]
    url = identity["homepage"].get("final_url") or identity["input"]["url"]
    m = identity["matches"]
    aliases = {norm_name(name)}
    if m["bioregistry"].get("status") == "found":
        aliases |= {norm_name(m["bioregistry"]["best"].get("name")), norm_name(m["bioregistry"]["best"]["id"])}
    if m["biotools"].get("status") == "found":
        aliases |= {norm_name(re.sub(r"\(.*?\)", "", m["biotools"]["record"].get("name", ""))), norm_name(m["biotools"]["id"])}
    aliases.discard("")

    out: dict = {"generated": now_iso(), "resource": {"name": name, "url": url, "aliases": sorted(aliases)}}

    # reference list freshness
    ages = {}
    for key in ("nodes", "cdr", "edd", "rir", "gcbr", "communities"):
        ts = (refs.get(key) or {}).get("retrieved")
        if ts:
            ages[key] = (datetime.now(timezone.utc) - datetime.fromisoformat(ts)).days
    out["reference_list_age_days"] = ages
    out["reference_lists_stale"] = any(a > 180 for a in ages.values()) or not refs

    # S1.4 existing status
    status = {}
    for key, label in (("cdr", "Core Data Resource (CDR)"), ("edd", "ELIXIR Deposition Database (EDD)"),
                       ("rir", "Recommended Interoperability Resource (RIR)"), ("gcbr", "Global Core Biodata Resource (GCBR)")):
        hits = [{"name": it["name"], "url": it.get("url"), "match_basis": b}
                for it in (refs.get(key) or {}).get("items", []) if (b := matches(it, name, url, aliases))]
        status[key] = {"label": label, "matched": bool(hits), "hits": hits,
                       "source_url": (refs.get(key) or {}).get("source_url")}
    out["status_match"] = status
    out["holds_existing_status"] = [v["label"] for v in status.values() if any(
        "homepage_domain" in h["match_basis"] or "name" in h["match_basis"] for h in v["hits"])]

    # S1.2 SDP listing + Node
    sdp_hits = []
    for node in (refs.get("nodes") or {}).get("items", []):
        for svc in node.get("services", []):
            b = matches(svc, name, url, aliases)
            if b:
                sdp_hits.append({"node": node["label"], "service_name": svc["name"], "service_url": svc.get("url"),
                                 "match_basis": b, "node_page": node["url"]})
    xlsx_hits = []
    if args.services_xlsx:
        try:
            for svc in load_xlsx_services(args.services_xlsx):
                b = matches(svc, name, url, aliases)
                if b:
                    xlsx_hits.append({**svc, "match_basis": b})
        except Exception as exc:  # noqa: BLE001
            out["services_xlsx_error"] = f"{type(exc).__name__}: {exc}"
    strong = [h for h in sdp_hits if "homepage_domain" in h["match_basis"] or "name" in h["match_basis"]]
    out["sdp_match"] = {"listed": bool(strong), "hits": sdp_hits[:10], "xlsx_hits": xlsx_hits[:10],
                        "nodes": sorted({h["node"] for h in strong} |
                                        {re.sub(r"^ELIXIR\s+", "ELIXIR-", str(h["node"])) for h in xlsx_hits
                                         if h.get("node") and ("homepage_domain" in h["match_basis"] or "name" in h["match_basis"])}),
                        "source_url": (refs.get("nodes") or {}).get("source_url")}

    # S1.1 node hints from affiliations
    hints = {}
    re3 = (load_json(ev / "registry_summary.json").get("re3data") or {}) if (ev / "registry_summary.json").exists() else {}
    for inst in re3.get("institutions") or []:
        node = ISO3_TO_NODE.get((inst.get("institutionCountry") or "").upper())
        if node:
            hints.setdefault(node, []).append(f"re3data institution: {inst.get('institutionName')}")
    epmc = load_json(ev / "registry_europepmc.json") if (ev / "registry_europepmc.json").exists() else {}
    for aff, _count in epmc.get("top_affiliations") or []:
        low = f" {aff.lower()}"
        for key, node in COUNTRY_TO_NODE.items():
            if key in low:
                hints.setdefault(node, []).append(f"recent database paper affiliation: {aff[:140]}")
    out["node_match"] = {"from_sdp": out["sdp_match"]["nodes"],
                         "affiliation_hints": {k: v[:5] for k, v in sorted(hints.items(), key=lambda kv: -len(kv[1]))},
                         "note": "S1.1 concerns the APPLICANT's Node membership; only the applicant can confirm it."}

    # S1.3 resource type signals
    signals = []
    for t in re3.get("types") or []:
        signals.append(f"re3data repository type: {t}")
    for u in re3.get("data_upload") or []:
        signals.append(f"re3data data upload: {u.get('dataUploadType')}")
    bt = m["biotools"].get("record") or {}
    if bt:
        signals.append(f"bio.tools toolType: {', '.join(bt.get('toolType') or [])}")
    desc = " ".join(filter(None, [re3.get("description"), bt.get("description"),
                                  (m["bioregistry"].get("best") or {}).get("description")])).lower()
    for kind, pat in (("curated knowledgebase", r"curat|literature|manually|expert"),
                      ("deposition database", r"deposit|submi(t|ssion)|archive|repository"),
                      ("biodata registry", r"registry|catalog(ue)?|index of"),
                      ("aggregation database", r"aggregat|integrat|compil|collect(s|ion) (data )?from")):
        hits = re.findall(pat, desc)
        if hits:
            signals.append(f"description keywords suggest {kind}: {sorted(set(hits))[:5]}")
    out["resource_type_signals"] = signals

    # S1.5 federation signals
    insts = re3.get("institutions") or []
    out["federation_signals"] = {"re3data_institution_count": len(insts),
                                 "institution_countries": sorted({i.get("institutionCountry") for i in insts if i.get("institutionCountry")}),
                                 "note": "Multiple institutions alone do not make a federated/consortium database; "
                                         "look for multi-site deployments, mirrors or joint branding."}

    # APP.community suggestions
    text = " ".join([desc, " ".join(bt and [t["term"] for t in bt.get("topic", [])] or []),
                     " ".join(re3.get("keywords") or []), " ".join(re3.get("subjects") or [])]).lower()
    web = load_json(ev / "website.json") if (ev / "website.json").exists() else {}
    elixir_links = (web.get("registry_links") or {}).get("elixir-europe.org", [])
    scores = []
    for comm in (refs.get("communities") or {}).get("items", []):
        pat = COMMUNITY_KEYWORDS.get(comm["name"])
        n = len(re.findall(pat, text)) if pat else 0
        linked = any(comm["url"].rstrip("/").rsplit("/", 1)[-1] in l for l in elixir_links)
        if n or linked:
            scores.append({"community": comm["name"], "keyword_hits": n, "linked_from_resource_website": linked,
                           "url": comm["url"]})
    out["community_suggestions"] = sorted(scores, key=lambda s: (s["linked_from_resource_website"], s["keyword_hits"]), reverse=True)[:5]
    out["communities_available"] = [c["name"] for c in (refs.get("communities") or {}).get("items", [])]

    emit(out, args.out, summary={"holds_status": out["holds_existing_status"], "sdp_listed": out["sdp_match"]["listed"],
                                 "sdp_nodes": out["sdp_match"]["nodes"],
                                 "community_top": [c["community"] for c in out["community_suggestions"][:3]],
                                 "stale_refs": out["reference_lists_stale"]})


if __name__ == "__main__":
    main()
