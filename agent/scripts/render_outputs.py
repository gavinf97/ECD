#!/usr/bin/env python3
"""
render_outputs.py — Phase 7 of the ECD Agent Skill.

Renders answers.json into the completed ECD Checklist, in the official
Checklist order (applicant details, Sections 1-8, Appendix 1 if reapplying):

  ECD_<slug>_checklist.json   copy of the answers document (+ validation report)
  ECD_<slug>_checklist.md     Markdown report: answers + status/confidence + evidence
  ECD_<slug>_checklist.html   self-contained HTML report (status chips, collapsible evidence, light/dark)
  ECD_<slug>_checklist.docx   Word form mirroring the Checklist layout (+ evidence appendix)
  ECD_<slug>_checklist.pdf    PDF of the DOCX (LibreOffice), reportlab fallback

--submission-copy renders DOCX/PDF without pre-fill statuses and without the evidence appendix,
i.e. what the applicant would email to elixir-community-databases@elixir-europe.org — only
sensible once `answers.py validate --strict` passes.

Usage:
    python render_outputs.py --answers WORKDIR/answers.json --outdir WORKDIR/output
                             [--formats json,md,html,docx,pdf] [--evidence-dir WORKDIR/evidence] [--submission-copy]
"""

from __future__ import annotations

import argparse
import base64
import json
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import SKILL_ROOT, load_json, load_schema, log, now_iso, slugify, write_json  # noqa: E402
from answers import option_values, validate  # noqa: E402

STATUS_LABEL = {"verified": "Verified", "inferred": "Inferred – review", "needs_applicant": "Applicant input needed",
                "applicant_provided": "Applicant provided", "not_found": "Not found", "query_failed": "Lookup failed",
                "pending": "Not assessed"}
PLACEHOLDER = "[Applicant to complete]"


def part_needed(spec: dict | None, choice) -> bool:
    """A part must be shown (with a placeholder if empty) when required, or when its required_if matches the choice."""
    spec = spec or {}
    return bool(spec.get("required")) or (spec.get("required_if") is not None and choice is not None
                                           and spec.get("required_if") == choice)


# ------------------------------------------------------------------ view model
def build_view(schema: dict, doc: dict, evidence_dir: Path | None) -> dict:
    answers = doc["answers"]
    reapp = doc["meta"].get("reapplication")
    sections = []
    for section in schema["sections"]:
        if section.get("conditional_on") == "reapplication" and not reapp:
            continue
        qs = []
        for q in section["questions"]:
            rec = answers.get(q["id"], {})
            parts = q["parts"]
            unresolved = rec.get("status") in ("pending", "needs_applicant", "query_failed")
            image = rec.get("image")
            image_data = None
            if image and Path(image).exists():
                image_data = "data:image/png;base64," + base64.b64encode(Path(image).read_bytes()).decode()
            qs.append({
                "id": q["id"], "number": q.get("number", ""), "group": q.get("group"), "subgroup": q.get("subgroup"),
                "prompt": q["prompt"], "in_form": q.get("in_form", True), "parts": parts,
                "choice_label": (parts.get("choice") or {}).get("label"),
                "options": option_values(schema, parts["choice"]["option_set"]) if "choice" in parts else None,
                "text_label": (parts.get("text") or {}).get("label"), "word_limit": (parts.get("text") or {}).get("word_limit"),
                "list_label": (parts.get("list") or {}).get("label"), "links_label": (parts.get("links") or {}).get("label"),
                "choice": rec.get("choice"), "text": rec.get("text"), "list": rec.get("list") or [],
                "links": rec.get("links") or [], "image": image, "image_data": image_data,
                "status": rec.get("status", "pending"), "status_label": STATUS_LABEL.get(rec.get("status"), rec.get("status")),
                "confidence": rec.get("confidence"), "evidence": rec.get("evidence") or [], "hint": rec.get("hint"),
                "notes": rec.get("notes"), "unresolved": unresolved,
                "show": {part: part_needed(spec, rec.get("choice")) for part, spec in parts.items()},
            })
        sections.append({"id": section["id"], "title": section["title"], "intro": section.get("intro"), "questions": qs})
    fair = None
    if evidence_dir and (evidence_dir / "fairchecker.json").exists():
        f = load_json(evidence_dir / "fairchecker.json")
        fair = {"scores": f.get("metric_scores") or [], "assessment_url": f.get("assessment_url"),
                "overall_percent": f.get("overall_percent"), "legend": f.get("score_legend")}
    return {"meta": doc["meta"], "schema": schema, "sections": sections, "fair": fair,
            "validation": validate(doc), "rendered": now_iso()}


# ------------------------------------------------------------------ markdown / html
def jinja_env():
    from jinja2 import Environment, FileSystemLoader, select_autoescape
    env = Environment(loader=FileSystemLoader(str(SKILL_ROOT / "templates")),
                      autoescape=select_autoescape(["html", "j2"], default_for_string=False),
                      trim_blocks=True, lstrip_blocks=True)
    return env


def render_md(view: dict, path: Path) -> None:
    env = jinja_env()
    env.autoescape = False
    path.write_text(env.get_template("checklist.md.j2").render(**view, placeholder=PLACEHOLDER), encoding="utf-8")


def render_html(view: dict, path: Path) -> None:
    path.write_text(jinja_env().get_template("checklist.html.j2").render(**view, placeholder=PLACEHOLDER), encoding="utf-8")


# ------------------------------------------------------------------ docx
def add_hyperlink(paragraph, url: str, text: str | None = None):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FA8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text or url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def render_docx(view: dict, path: Path, submission: bool) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(view["schema"]["title"])
    r.bold, r.font.size = True, Pt(20)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Checklist {view['schema']['checklist_version']}").font.size = Pt(12)
    src = doc.add_paragraph()
    src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    src.add_run("Checklist source: ").font.size = Pt(9)
    add_hyperlink(src, f"https://doi.org/{view['schema']['source']['checklist']['doi']}")
    if not submission:
        note = doc.add_paragraph()
        nr = note.add_run("Pre-filled by the ECD Agent Skill from public registries and the resource website. Every answer "
                          "must be reviewed by the applicant before submission. Highlighted items still need applicant input.")
        nr.italic, nr.font.size = True, Pt(9)
        nr.font.color.rgb = RGBColor(0x80, 0x4A, 0x00)

    def answer_line(par_text: str | None, label: str | None = None, highlight=False, bold=False):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.9)
        if label:
            lr = p.add_run(f"{label}: ")
            lr.font.size = Pt(9)
            lr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        run = p.add_run(par_text or PLACEHOLDER)
        run.bold = bold
        if highlight or not par_text:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return p

    def status_suffix(p, q):
        if submission:
            return
        sr = p.add_run(f"   [{q['status_label']}{' · ' + q['confidence'] if q['confidence'] else ''}]")
        sr.font.size = Pt(8)
        sr.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)

    for section in view["sections"]:
        if section["id"] == "APP":
            doc.add_heading("Applicant details", level=1)
            table = doc.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for q in section["questions"]:
                row = table.add_row().cells
                row[0].text = q["prompt"]
                val = q["choice"] or q["text"] or ", ".join(q["links"]) or ""
                para = row[1].paragraphs[0]
                if q["links"] and not q["choice"] and not q["text"]:
                    for i, link in enumerate(q["links"]):
                        if i:
                            para.add_run(", ")
                        add_hyperlink(para, link)
                else:
                    run = para.add_run(val or PLACEHOLDER)
                    if not val:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                if not submission:
                    sr = para.add_run(f"  [{q['status_label']}]")
                    sr.font.size = Pt(7)
            continue

        doc.add_heading(section["title"], level=1)
        if section.get("intro"):
            doc.add_paragraph(section["intro"]).runs[0].italic = True
        last_group = last_sub = None
        for q in section["questions"]:
            if q.get("subgroup") and q["subgroup"] != last_sub:
                doc.add_heading(q["subgroup"], level=2)
                last_sub = q["subgroup"]
            if q.get("group") and q["group"] != last_group:
                gp = doc.add_paragraph()
                gp.add_run(q["group"]).bold = True
                last_group = q["group"]
            qp = doc.add_paragraph()
            qp.paragraph_format.space_before = Pt(6)
            qr = qp.add_run(f"{q['number']} {q['prompt']}".strip())
            qr.bold = q["in_form"]
            if not q["in_form"]:
                if submission:
                    continue
                qr.italic = True
            if q["choice"] is not None or "choice" in q["parts"]:
                p = answer_line(q["choice"], q["choice_label"] or "Answer", highlight=q["unresolved"], bold=True)
                status_suffix(p, q)
            if "list" in q["parts"]:
                if not q["list"] and part_needed(q["parts"]["list"], q["choice"]):
                    answer_line(None, q["list_label"] or "List")
                if q["list"]:
                    if q["list_label"]:
                        lp = doc.add_paragraph()
                        lp.paragraph_format.left_indent = Cm(0.9)
                        lr = lp.add_run(q["list_label"])
                        lr.font.size = Pt(9)
                    for item in q["list"]:
                        bp = doc.add_paragraph(str(item), style="List Bullet")
                        bp.paragraph_format.left_indent = Cm(1.6)
            if "text" in q["parts"]:
                if q["text"] or part_needed(q["parts"]["text"], q["choice"]):
                    label = q["text_label"] or "Response"
                    if q["word_limit"]:
                        label += f" [{q['word_limit']} words]"
                    answer_line(q["text"], label, highlight=not q["text"])
            if "links" in q["parts"] and (q["links"] or part_needed(q["parts"]["links"], q["choice"])):
                lp = doc.add_paragraph()
                lp.paragraph_format.left_indent = Cm(0.9)
                lab = lp.add_run(f"{q['links_label'] or 'Link(s)'}: ")
                lab.font.size = Pt(9)
                if q["links"]:
                    for i, link in enumerate(q["links"]):
                        if i:
                            lp.add_run("  ·  ")
                        add_hyperlink(lp, link)
                else:
                    lp.add_run(PLACEHOLDER).font.highlight_color = WD_COLOR_INDEX.YELLOW
            if "image" in q["parts"]:
                if q["image"] and Path(q["image"]).exists():
                    doc.add_picture(q["image"], width=Cm(15.5))
                    if view.get("fair") and view["fair"].get("assessment_url"):
                        fp = doc.add_paragraph()
                        fp.add_run("FAIR-Checker assessment: ").font.size = Pt(9)
                        add_hyperlink(fp, view["fair"]["assessment_url"])
                else:
                    answer_line(None, "Screenshot", highlight=True)
            if not submission and q["hint"] and q["unresolved"]:
                hp = doc.add_paragraph()
                hp.paragraph_format.left_indent = Cm(0.9)
                hr = hp.add_run(f"Hint: {q['hint']}")
                hr.italic, hr.font.size = True, Pt(8.5)

    if not submission:
        doc.add_page_break()
        doc.add_heading("Evidence appendix (not part of the submission)", level=1)
        doc.add_paragraph(f"Generated {view['rendered']} by the ECD Agent Skill. Status counts: "
                          + ", ".join(f"{STATUS_LABEL.get(k, k)}: {v}" for k, v in view["validation"]["status_counts"].items()))
        for section in view["sections"]:
            for q in section["questions"]:
                if not q["evidence"]:
                    continue
                ep = doc.add_paragraph()
                er = ep.add_run(f"{q['id']} — {q['status_label']}{' (' + q['confidence'] + ')' if q['confidence'] else ''}")
                er.bold, er.font.size = True, Pt(9)
                for e in q["evidence"][:6]:
                    bp = doc.add_paragraph(style="List Bullet")
                    br = bp.add_run(f"{e.get('source')}: ")
                    br.font.size = Pt(8.5)
                    if e.get("url"):
                        add_hyperlink(bp, e["url"])
                    if e.get("snippet"):
                        sr = bp.add_run(f" — “{e['snippet'][:240]}”")
                        sr.font.size = Pt(8.5)
    doc.save(path)


def render_pdf(docx_path: Path, pdf_path: Path, view: dict) -> str:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
                           capture_output=True, timeout=180, check=True)
            produced = pdf_path.parent / (docx_path.stem + ".pdf")
            if produced != pdf_path and produced.exists():
                produced.replace(pdf_path)
            if pdf_path.exists():
                return "libreoffice"
        except Exception as exc:  # noqa: BLE001
            log(f"[render] LibreOffice conversion failed: {exc}; falling back to reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(view["schema"]["title"]), styles["Title"])]
    for section in view["sections"]:
        story.append(Paragraph(escape(section["title"]), styles["Heading2"]))
        for q in section["questions"]:
            story.append(Paragraph(f"<b>{escape(q['number'])} {escape(q['prompt'])}</b>", styles["Normal"]))
            val = " | ".join(filter(None, [q["choice"], q["text"], "; ".join(map(str, q["list"])), " ".join(q["links"])]))
            story.append(Paragraph(escape(val or PLACEHOLDER), styles["Normal"]))
            story.append(Spacer(1, 4))
    SimpleDocTemplate(str(pdf_path), pagesize=A4).build(story)
    return "reportlab"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--answers", required=True)
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--formats", default="json,md,html,docx,pdf")
    ap.add_argument("--evidence-dir")
    ap.add_argument("--submission-copy", action="store_true")
    args = ap.parse_args()

    schema = load_schema()
    doc = load_json(args.answers)
    evidence_dir = Path(args.evidence_dir) if args.evidence_dir else Path(args.answers).parent / "evidence"
    view = build_view(schema, doc, evidence_dir if evidence_dir.exists() else None)
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    stem = f"ECD_{slugify(doc['meta']['resource_name'])}_checklist" + ("_submission" if args.submission_copy else "")
    formats = [f.strip() for f in args.formats.split(",") if f.strip()]
    written = {}
    if "json" in formats:
        write_json(outdir / f"{stem}.json", {**doc, "validation": view["validation"]})
        written["json"] = str(outdir / f"{stem}.json")
    if "md" in formats:
        render_md(view, outdir / f"{stem}.md")
        written["md"] = str(outdir / f"{stem}.md")
    if "html" in formats:
        render_html(view, outdir / f"{stem}.html")
        written["html"] = str(outdir / f"{stem}.html")
    if "docx" in formats or "pdf" in formats:
        render_docx(view, outdir / f"{stem}.docx", args.submission_copy)
        written["docx"] = str(outdir / f"{stem}.docx")
    if "pdf" in formats:
        engine = render_pdf(outdir / f"{stem}.docx", outdir / f"{stem}.pdf", view)
        written["pdf"] = str(outdir / f"{stem}.pdf")
        written["pdf_engine"] = engine
    print(json.dumps({"written": written, "validation": {k: view["validation"][k] for k in ("valid", "status_counts", "ready_to_submit")},
                      "unresolved": view["validation"]["unresolved"]}, indent=2))


if __name__ == "__main__":
    main()
